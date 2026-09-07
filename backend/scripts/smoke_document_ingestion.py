"""Exercise DOCX parsing, ingestion and persistence without external services."""

from __future__ import annotations

import asyncio
import os
import tempfile
from pathlib import Path

os.environ["DATABASE_URL"] = "sqlite+aiosqlite://"
os.environ["EMBEDDING_PROVIDER"] = "hashing"

from docx import Document
from sqlalchemy import select

from app.db import KnowledgeChunkRow, KnowledgeDocumentRow, SessionLocal, create_schema
from app.ingestion import _persist


async def main() -> None:
    await create_schema()
    with tempfile.TemporaryDirectory(prefix="ielts-ingestion-") as directory:
        path = Path(directory) / "project-original-study-note.docx"
        document = Document()
        document.add_heading("Project-authored IELTS study note", level=1)
        document.add_paragraph(
            "A clear paragraph develops one main idea with an explanation and a relevant example."
        )
        document.save(path)

        async with SessionLocal() as session:
            row = KnowledgeDocumentRow(
                title="DOCX smoke note",
                filename=path.name,
                license_type="project-original",
                status="queued",
            )
            session.add(row)
            await session.commit()
            await session.refresh(row)
            document_id = row.id

        await _persist(document_id, path)

        async with SessionLocal() as session:
            stored = await session.get(KnowledgeDocumentRow, document_id)
            chunks = (
                await session.scalars(
                    select(KnowledgeChunkRow).where(
                        KnowledgeChunkRow.document_id == document_id
                    )
                )
            ).all()
            if stored is None or stored.status != "completed" or not chunks:
                raise RuntimeError(
                    f"DOCX ingestion failed: status={getattr(stored, 'status', None)}, "
                    f"error={getattr(stored, 'error', None)}"
                )
            print(f"status {stored.status}")
            print(f"chunks {len(chunks)}")
            print(f"embedding_dimensions {len(chunks[0].embedding or [])}")
            print("docx_ingestion_valid True")


if __name__ == "__main__":
    asyncio.run(main())
